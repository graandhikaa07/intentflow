import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation('docs/Review2_UserIntent_LLM.pptx')
print(f"Slide Dimensions: {prs.slide_width/914400:.3f} x {prs.slide_height/914400:.3f} inches")
print(f"Total Slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides, 1):
    print("=" * 80)
    print(f"SLIDE {i}: (Layout: {slide.slide_layout.name if slide.slide_layout else 'None'})")
    for s_idx, shape in enumerate(slide.shapes):
        stype = shape.shape_type
        name = shape.name
        left, top, w, h = shape.left/914400, shape.top/914400, shape.width/914400, shape.height/914400
        
        fill_info = "None"
        try:
            if shape.fill.type:
                fill_info = f"Type: {shape.fill.type}"
                if hasattr(shape.fill, 'fore_color') and shape.fill.fore_color:
                    if shape.fill.fore_color.type:
                        fill_info += f", Color: {shape.fill.fore_color.rgb if hasattr(shape.fill.fore_color, 'rgb') else shape.fill.fore_color.theme_color}"
        except:
            pass
            
        print(f"  Shape {s_idx} [{name}] ({stype}): pos=({left:.2f}, {top:.2f}), size=({w:.2f} x {h:.2f}), fill={fill_info}")
        
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    font_info = ""
                    if p.font:
                        fn = p.font.name
                        fs = p.font.size.pt if p.font.size else "auto"
                        b = "B" if p.font.bold else ""
                        c = ""
                        try:
                            if p.font.color and hasattr(p.font.color, 'rgb') and p.font.color.rgb:
                                c = f"#{p.font.color.rgb}"
                        except:
                            pass
                        font_info = f"[{fn} {fs}pt {b} {c}]"
                    print(f"    Text: {font_info} {txt[:100]}")
                    
        if shape.has_table:
            tbl = shape.table
            print(f"    TABLE: {len(tbl.rows)} rows x {len(tbl.columns)} cols")
            for r_idx, row in enumerate(tbl.rows[:3]):
                cells_txt = [c.text.strip().replace('\n', ' ')[:30] for c in row.cells]
                print(f"      Row {r_idx}: {' | '.join(cells_txt)}")
