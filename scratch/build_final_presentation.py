"""
Build Final 10-Slide Academic Project Presentation (Revised)
Incorporates:
- 1 Architecture Diagram (Slide 4)
- 1 Confusion Matrix (Slide 5)
- 2 Result Graphs: Efficiency Comparison (Slide 7) & Per-Intent Accuracy (Slide 8)
- Drastically reduced text on every slide for clean, high-impact readability
- Preserves the visual identity, typography, color palette, and layout of docs/Review2_UserIntent_LLM.pptx
Outputs to artifacts/reports/FINAL_PROJECT_PRESENTATION.pptx
"""

import os
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PROJECT_ROOT = Path(__file__).resolve().parent.parent

# Define Palette from original presentation
COLOR_NAVY_DARK = RGBColor(0x1A, 0x1A, 0x2E)    # #1A1A2E (Header bar / Title background)
COLOR_NAVY_BLUE = RGBColor(0x1F, 0x3A, 0x5F)    # #1F3A5F (Primary Accent / Card Header)
COLOR_SLATE_DARK = RGBColor(0x2C, 0x3E, 0x50)   # #2C3E50 (Secondary Accent)
COLOR_TEXT_DARK = RGBColor(0x1E, 0x29, 0x3B)    # #1E293B (Primary dark text)
COLOR_TEXT_CHARCOAL = RGBColor(0x2C, 0x2C, 0x2C)# #2C2C2C (Body text)
COLOR_TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)   # #64748B (Secondary muted text)
COLOR_TEXT_LIGHT = RGBColor(0xAA, 0xBB, 0xD0)   # #AABBD0 (Title slide subtitle)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)        # #FFFFFF (White)
COLOR_CARD_BG = RGBColor(0xF2, 0xF2, 0xF2)      # #F2F2F2 (Soft grey card fill)
COLOR_CARD_WHITE = RGBColor(0xFF, 0xFF, 0xFF)   # #FFFFFF
COLOR_BORDER_GREY = RGBColor(0xCB, 0xD5, 0xE1)  # #CBD5E1 (Light grey border)
COLOR_ACCENT_BLUE = RGBColor(0x02, 0x84, 0xC7)  # #0284C7 (Vibrant cyan-blue highlight)
COLOR_SUCCESS_GREEN = RGBColor(0x10, 0xB9, 0x81)# #10B981 (Success / metric highlight)
COLOR_AMBER = RGBColor(0xD9, 0x77, 0x06)        # #D97706 (Tradeoff amber)

FONT_FAMILY = "Calibri"

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(10.0)
    prs.slide_height = Inches(5.625)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    # Helper: Add Standard Header Bar and Footer
    def add_standard_chrome(slide, title_text, slide_num):
        # Header banner
        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(0.0), Inches(10.0), Inches(0.65)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = COLOR_NAVY_DARK
        header_shape.line.fill.background()

        # Title text
        txBox = slide.shapes.add_textbox(Inches(0.40), Inches(0.08), Inches(8.50), Inches(0.50))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

        # Slide Number Footer
        ftBox = slide.shapes.add_textbox(Inches(9.10), Inches(5.35), Inches(0.70), Inches(0.20))
        ft_tf = ftBox.text_frame
        ft_tf.margin_left = ft_tf.margin_top = ft_tf.margin_right = ft_tf.margin_bottom = 0
        p_ft = ft_tf.paragraphs[0]
        p_ft.alignment = PP_ALIGN.RIGHT
        p_ft.text = f"Slide {slide_num}"
        p_ft.font.name = FONT_FAMILY
        p_ft.font.size = Pt(8)
        p_ft.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Helper: Add Card Box
    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER_GREY):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # Helper: Add Card with Top Header Strip
    def add_header_card(slide, left, top, width, height, header_title, header_bg=COLOR_NAVY_BLUE, card_bg=COLOR_CARD_BG, header_h=Inches(0.36)):
        add_card(slide, left, top, width, height, bg_color=card_bg, border_color=COLOR_BORDER_GREY)
        strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, header_h)
        strip.fill.solid()
        strip.fill.fore_color.rgb = header_bg
        strip.line.fill.background()

        txBox = slide.shapes.add_textbox(left + Inches(0.12), top + Inches(0.04), width - Inches(0.24), header_h - Inches(0.08))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = header_title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.LEFT

    # Helper: Add Stat Metric Card
    def add_stat_card(slide, left, top, width, height, label, main_val, sub_val, val_color=COLOR_NAVY_BLUE):
        add_header_card(slide, left, top, width, height, label, header_bg=COLOR_SLATE_DARK, card_bg=COLOR_CARD_BG, header_h=Inches(0.30))
        tb = slide.shapes.add_textbox(left + Inches(0.08), top + Inches(0.35), width - Inches(0.16), height - Inches(0.40))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p_main = tf.paragraphs[0]
        p_main.alignment = PP_ALIGN.CENTER
        p_main.text = main_val
        p_main.font.name = FONT_FAMILY
        p_main.font.size = Pt(22)
        p_main.font.bold = True
        p_main.font.color.rgb = val_color

        p_sub = tf.add_paragraph()
        p_sub.alignment = PP_ALIGN.CENTER
        p_sub.text = sub_val
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(8)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Helper: Set Speaker Notes
    def set_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text.strip()

    # =========================================================================
    # SLIDE 1 — TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1_top = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(0.0), Inches(10.0), Inches(3.60))
    bg1_top.fill.solid()
    bg1_top.fill.fore_color.rgb = COLOR_NAVY_DARK
    bg1_top.line.fill.background()

    bg1_bot = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(3.60), Inches(10.0), Inches(2.025))
    bg1_bot.fill.solid()
    bg1_bot.fill.fore_color.rgb = COLOR_WHITE
    bg1_bot.line.fill.background()

    line1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(3.57), Inches(10.0), Inches(0.03))
    line1.fill.solid()
    line1.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    line1.line.fill.background()

    tb_t = slide1.shapes.add_textbox(Inches(0.60), Inches(0.75), Inches(8.80), Inches(1.80))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0

    p_t = tf_t.paragraphs[0]
    p_t.text = "User Intent Classification for\nHallucination Reduction in LLMs"
    p_t.font.name = FONT_FAMILY
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_WHITE

    p_sub = tf_t.add_paragraph()
    p_sub.space_before = Pt(8)
    p_sub.text = "A Cost–Accuracy–Latency Tradeoff Study"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_TEXT_LIGHT

    p_framing = tf_t.add_paragraph()
    p_framing.space_before = Pt(4)
    p_framing.text = "Pre-Generation Query Routing Across Deterministic Symbolic & Foundation LLM Engines"
    p_framing.font.name = FONT_FAMILY
    p_framing.font.size = Pt(11)
    p_framing.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    tb_a = slide1.shapes.add_textbox(Inches(0.60), Inches(3.90), Inches(8.80), Inches(1.20))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    tf_a.margin_left = tf_a.margin_top = tf_a.margin_right = tf_a.margin_bottom = 0

    p_auth = tf_a.paragraphs[0]
    p_auth.text = "P.U. Graandhikaa Sri (CH.SC.U4CSE23136)   |   T. L. Bharath (CH.SC.U4CSE23149)"
    p_auth.font.name = FONT_FAMILY
    p_auth.font.size = Pt(13)
    p_auth.font.bold = True
    p_auth.font.color.rgb = COLOR_TEXT_DARK

    p_dt = tf_a.add_paragraph()
    p_dt.space_before = Pt(4)
    p_dt.text = "Academic Research Project  •  Department of Computer Science & Engineering"
    p_dt.font.name = FONT_FAMILY
    p_dt.font.size = Pt(10.5)
    p_dt.font.color.rgb = COLOR_TEXT_MUTED

    set_notes(slide1, (
        "Good morning, esteemed professors and members of the review panel. "
        "Today we present our research on User Intent Classification for Hallucination Reduction in LLMs: "
        "A Cost-Accuracy Tradeoff Study. "
        "Our work explores whether classifying user queries before generation enables intelligent routing "
        "to specialized deterministic engines and cloud LLMs, establishing a practical Pareto frontier."
    ))

    # =========================================================================
    # SLIDE 2 — PROBLEM & MOTIVATION (LOW TEXT)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide2, "Problem Formulation & Research Motivation", 2)

    # Top Central Problem Card
    add_card(slide2, Inches(0.40), Inches(0.80), Inches(9.20), Inches(0.65), bg_color=COLOR_WHITE, border_color=COLOR_ACCENT_BLUE)
    tb_pr = slide2.shapes.add_textbox(Inches(0.55), Inches(0.85), Inches(8.90), Inches(0.55))
    tf_pr = tb_pr.text_frame
    tf_pr.word_wrap = True
    tf_pr.margin_left = tf_pr.margin_top = tf_pr.margin_right = tf_pr.margin_bottom = 0
    p = tf_pr.paragraphs[0]
    p.text = "The Monolithic Dilemma in Production AI:"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_BLUE
    p_sub = tf_pr.add_paragraph()
    p_sub.text = "Treating every user query as an autoregressive text generation problem causes hallucinations, high latency, and excessive cloud API spend."
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = COLOR_TEXT_CHARCOAL

    # 4 Archetype Cards across
    card_w = Inches(2.15)
    card_h = Inches(2.80)
    gap_x = Inches(0.20)
    lefts = [Inches(0.40) + i * (card_w + gap_x) for i in range(4)]

    archetypes = [
        ("Arithmetic & Math", COLOR_NAVY_BLUE, "🧮",
         "Computation",
         "• LLMs predict math tokens probabilistically\n• Frequent arithmetic & calculation slips\n• Solvable deterministically via symbolic engines (SymPy)"),

        ("Factual Inquiries", COLOR_SLATE_DARK, "📚",
         "Attribution",
         "• Parametric memory is lossy and decays\n• Hallucinates plausible dates & citations\n• Requires verified grounding in external reference corpus"),

        ("Creative & Style", COLOR_NAVY_BLUE, "🎨",
         "Adaptation",
         "• Subjective tone, persona & style shifts\n• Symbolic code cannot adapt creative tone\n• Foundation LLMs (Gemini) excel here"),

        ("Structured Schema", COLOR_SLATE_DARK, "⚙️",
         "Serialization",
         "• Strict JSON/YAML formatting demands\n• LLMs emit conversational preambles\n• Best handled by programmatic schema solvers")
    ]

    for i, (title, color, icon, tag, desc) in enumerate(archetypes):
        add_header_card(slide2, lefts[i], Inches(1.60), card_w, card_h, f"{icon} {title}", header_bg=color)
        tb = slide2.shapes.add_textbox(lefts[i] + Inches(0.10), Inches(2.02), card_w - Inches(0.20), card_h - Inches(0.48))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p_tag = tf.paragraphs[0]
        p_tag.text = f"Needs: {tag}"
        p_tag.font.name = FONT_FAMILY
        p_tag.font.size = Pt(8.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_ACCENT_BLUE

        p_desc = tf.add_paragraph()
        p_desc.space_before = Pt(4)
        p_desc.text = desc
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(8)
        p_desc.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Bottom Takeaway Card
    add_card(slide2, Inches(0.40), Inches(4.55), Inches(9.20), Inches(0.60), bg_color=RGBColor(0xF0, 0xFD, 0xF4), border_color=COLOR_SUCCESS_GREEN)
    tb_bot = slide2.shapes.add_textbox(Inches(0.55), Inches(4.62), Inches(8.90), Inches(0.45))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    tf_bot.margin_left = tf_bot.margin_top = tf_bot.margin_right = tf_bot.margin_bottom = 0
    p = tf_bot.paragraphs[0]
    p.text = "Research Opportunity:"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
    p_sub = tf_bot.add_paragraph()
    p_sub.text = "Pre-generation intent classification allows routing deterministic queries away from expensive cloud LLMs to specialized zero-cost local tools."
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(8.5)
    p_sub.font.color.rgb = COLOR_TEXT_DARK

    set_notes(slide2, (
        "Modern LLM architectures treat all incoming user queries as autoregressive text completions. "
        "However, different query modalities have fundamentally different computational requirements. "
        "Math requires exact symbolic execution, factual queries require verified reference retrieval, "
        "schema formatting requires strict serialization, and only creative queries genuinely need cloud LLM generation. "
        "Treating all four through a single monolithic model incurs high costs, latency, and hallucinations."
    ))

    # =========================================================================
    # SLIDE 3 — PROPOSED APPROACH (LOW TEXT)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide3, "Proposed Approach: Pre-Generation Intent Routing", 3)

    # Core Concept Strip
    add_card(slide3, Inches(0.40), Inches(0.80), Inches(9.20), Inches(0.60), bg_color=COLOR_WHITE, border_color=COLOR_ACCENT_BLUE)
    tb_c = slide3.shapes.add_textbox(Inches(0.55), Inches(0.85), Inches(8.90), Inches(0.50))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
    p = tf_c.paragraphs[0]
    p.text = "Core Thesis: Intent-Aware Modality Dispatch"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_BLUE
    p_sub = tf_c.add_paragraph()
    p_sub.text = "Classify the query intent first using a lightweight encoder (<50 ms), then route directly to the optimal specialized execution venue."
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = COLOR_TEXT_CHARCOAL

    # 4 Intent Taxonomy Cards
    intents = [
        ("MATHDS", "Mathematical Reasoning", COLOR_NAVY_BLUE,
         "Local SymPy Solver",
         "• Exact symbolic arithmetic\n• 0ms API latency\n• Eliminates calculation hallucinations"),

        ("TXTSRC", "Factual Information", COLOR_SLATE_DARK,
         "Reference Corpus RAG",
         "• Grounded passage retrieval\n• Explicit [REF_01..20] citations\n• Eliminates ungrounded confabulation"),

        ("SEMADPT", "Adaptive & Creative", COLOR_NAVY_BLUE,
         "Google Gemini 3.1 Flash-Lite",
         "• Autoregressive cloud LLM\n• Creative tone, haiku & rewrites\n• Incurs cloud tokens only when needed"),

        ("CSTRSET", "Constrained Schema", COLOR_SLATE_DARK,
         "Deterministic Synthesizer",
         "• Programmatic JSON/YAML engine\n• Guaranteed structural compliance\n• Zero conversational syntax drift")
    ]

    for i, (code, title, color, engine, desc) in enumerate(intents):
        add_header_card(slide3, lefts[i], Inches(1.55), card_w, Inches(2.85), f"[{code}]", header_bg=color)
        tb = slide3.shapes.add_textbox(lefts[i] + Inches(0.10), Inches(1.95), card_w - Inches(0.20), Inches(2.40))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(9)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY_BLUE

        p_eng = tf.add_paragraph()
        p_eng.space_before = Pt(3)
        p_eng.text = f"Handler: {engine}"
        p_eng.font.name = FONT_FAMILY
        p_eng.font.size = Pt(8)
        p_eng.font.bold = True
        p_eng.font.color.rgb = COLOR_ACCENT_BLUE

        p_desc = tf.add_paragraph()
        p_desc.space_before = Pt(4)
        p_desc.text = desc
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(8)
        p_desc.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Bottom Callout Bar
    add_card(slide3, Inches(0.40), Inches(4.55), Inches(9.20), Inches(0.60), bg_color=RGBColor(0xF8, 0xFA, 0xFC), border_color=COLOR_BORDER_GREY)
    tb_b = slide3.shapes.add_textbox(Inches(0.55), Inches(4.62), Inches(8.90), Inches(0.45))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
    p = tf_b.paragraphs[0]
    p.text = "Key Operational Paradigm:"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_BLUE
    p_sub = tf_b.add_paragraph()
    p_sub.text = "3 of the 4 execution paths execute 100% locally on standard CPU hardware without network round-trips or cloud API tokens."
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(8.5)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    set_notes(slide3, (
        "Our proposed solution is pre-generation intent routing. "
        "We define a compact, four-intent operational taxonomy: "
        "MATHDS for deterministic arithmetic, TXTSRC for factual retrieval, "
        "SEMADPT for creative adaptation, and CSTRSET for structured schemas. "
        "By classifying the intent up front, we only call the cloud LLM when open-ended generation is truly needed, "
        "delegating all deterministic tasks to local engines."
    ))

    # =========================================================================
    # SLIDE 4 — SYSTEM ARCHITECTURE (WITH ARCHITECTURE DIAGRAM FIGURE)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide4, "System Architecture: IntentFlow Pipeline", 4)

    # Left: Architecture Diagram Figure (High-Res Image)
    arch_img_path = PROJECT_ROOT / "artifacts" / "figures" / "fig1_system_architecture.png"
    if arch_img_path.exists():
        slide4.shapes.add_picture(str(arch_img_path), Inches(0.40), Inches(0.80), width=Inches(6.30), height=Inches(3.90))

    # Right: Pipeline Architecture Highlights Card
    r_x = Inches(6.85)
    r_w = Inches(2.75)
    r_h = Inches(3.90)
    add_header_card(slide4, r_x, Inches(0.80), r_w, r_h, "Architectural Principles", header_bg=COLOR_NAVY_BLUE)
    tb_arch = slide4.shapes.add_textbox(r_x + Inches(0.12), Inches(1.22), r_w - Inches(0.24), r_h - Inches(0.48))
    tf_arch = tb_arch.text_frame
    tf_arch.word_wrap = True
    tf_arch.margin_left = tf_arch.margin_top = tf_arch.margin_right = tf_arch.margin_bottom = 0

    points = [
        ("Pre-Generation Gate", "Lightweight DistilBERT classifies incoming query in ~41 ms."),
        ("Autonomous Dispatch", "QueryRouter maps intent label to specialized domain handler."),
        ("77.5% Local Execution", "Math, RAG, and Schema engines run 100% locally on standard CPU."),
        ("Zero-Cost Offload", "Only creative SEMADPT prompts invoke the cloud Gemini API."),
        ("Telemetry Synthesis", "Aggregator logs exact latency, token costs, and provenance.")
    ]

    for idx, (hd, desc) in enumerate(points):
        p_h = tf_arch.paragraphs[0] if idx == 0 else tf_arch.add_paragraph()
        if idx > 0: p_h.space_before = Pt(3)
        p_h.text = f"• {hd}:"
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(8.5)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_NAVY_BLUE

        p_d = tf_arch.add_paragraph()
        p_d.text = f"  {desc}"
        p_d.font.name = FONT_FAMILY
        p_d.font.size = Pt(7.5)
        p_d.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Bottom Note
    add_card(slide4, Inches(0.40), Inches(4.80), Inches(9.20), Inches(0.45), bg_color=RGBColor(0xF8, 0xFA, 0xFC), border_color=COLOR_BORDER_GREY)
    tb_ui_nt = slide4.shapes.add_textbox(Inches(0.55), Inches(4.85), Inches(8.90), Inches(0.35))
    tf_ui_nt = tb_ui_nt.text_frame
    tf_ui_nt.word_wrap = True
    tf_ui_nt.margin_left = tf_ui_nt.margin_top = tf_ui_nt.margin_right = tf_ui_nt.margin_bottom = 0
    p = tf_ui_nt.paragraphs[0]
    p.text = "Deployment Note: Interactive UI supports both autonomous auto-detection and optional user intent override for testing."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE

    set_notes(slide4, (
        "Here is the complete architectural topology of IntentFlow. "
        "A user submits a query, which enters our pre-generation DistilBERT classifier. "
        "In approximately 41 milliseconds, it predicts the intent class. "
        "The QueryRouter immediately dispatches the query to SymPy for math, our reference corpus for factual inquiries, "
        "our programmatic schema synthesizer for structured output, or Google Gemini for creative tasks. "
        "Because three of the four routes execute locally, 77.5% of queries bypass the cloud LLM entirely."
    ))

    # =========================================================================
    # SLIDE 5 — IMPLEMENTATION & CONFUSION MATRIX (CONFUSION MATRIX FIGURE)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide5, "System Implementation & Classifier Performance", 5)

    # Left: 3 Implementation Specification Cards
    l_w = Inches(4.30)
    l_x = Inches(0.40)
    
    spec_cards = [
        ("DistilBERT Classifier Engine", COLOR_NAVY_BLUE, Inches(0.80), Inches(1.30),
         "• Architecture: DistilBERT-base-uncased (66M params)\n"
         "• Test Accuracy: 99.77% on held-out test split (N=2,650)\n"
         "• Latency: ~41 ms inference on standard CPU"),

        ("Curated Training Dataset (V3)", COLOR_SLATE_DARK, Inches(2.20), Inches(1.30),
         "• Scale: 26,492 total curated training records\n"
         "• Provenance: 11 public benchmarks (CLINC, GSM8K, TriviaQA, IFEval)\n"
         "• Zero-Leakage: Cryptographically partitioned from evaluation"),

        ("Adversarial Boundary Augmentation", COLOR_NAVY_BLUE, Inches(3.60), Inches(1.30),
         "• Targeted Augmentation: +800 paired boundary prompts\n"
         "• Objective: Explicitly hardens the creative vs. schema boundary\n"
         "• Resolves ambiguity between poetic constraints and JSON schemas")
    ]

    for title, col, y, h, text in spec_cards:
        add_header_card(slide5, l_x, y, l_w, h, title, header_bg=col, header_h=Inches(0.30))
        tb = slide5.shapes.add_textbox(l_x + Inches(0.10), y + Inches(0.34), l_w - Inches(0.20), h - Inches(0.38))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(8)
        p.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Right: Confusion Matrix Figure
    cm_img_path = PROJECT_ROOT / "artifacts" / "figures" / "fig5_confusion_matrix_final.png"
    if cm_img_path.exists():
        slide5.shapes.add_picture(str(cm_img_path), Inches(4.90), Inches(0.80), width=Inches(4.70), height=Inches(3.80))

    # Bottom Strip
    add_card(slide5, Inches(0.40), Inches(5.00), Inches(9.20), Inches(0.32), bg_color=RGBColor(0xF8, 0xFA, 0xFC), border_color=COLOR_BORDER_GREY)
    tb_cmnote = slide5.shapes.add_textbox(Inches(0.55), Inches(5.03), Inches(8.90), Inches(0.26))
    tf_cmnote = tb_cmnote.text_frame
    tf_cmnote.word_wrap = True
    tf_cmnote.margin_left = tf_cmnote.margin_top = tf_cmnote.margin_right = tf_cmnote.margin_bottom = 0
    p = tf_cmnote.paragraphs[0]
    p.text = "Empirical Validation: DistilBERT correctly classified 2,644 of 2,650 test samples (99.77% accuracy, Macro F1: 99.77%)."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_BLUE

    set_notes(slide5, (
        "To ensure the router introduces negligible latency, we fine-tuned a compact 66-million parameter DistilBERT model. "
        "As shown in the confusion matrix on the right, the model achieves 99.77% accuracy on our 2,650-sample held-out test split, "
        "making only 6 errors across the entire test set. "
        "We also augmented the dataset with 800 adversarial boundary samples to specifically address ambiguity "
        "between creative constraints like rhyming poems and structural schema constraints."
    ))

    # =========================================================================
    # SLIDE 6 — EXPERIMENTAL DESIGN (LOW TEXT)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide6, "Experimental Design & Matched-Pair Evaluation", 6)

    # Top Benchmark Banner
    add_card(slide6, Inches(0.40), Inches(0.80), Inches(9.20), Inches(0.80), bg_color=COLOR_WHITE, border_color=COLOR_ACCENT_BLUE)
    tb_bn = slide6.shapes.add_textbox(Inches(0.55), Inches(0.85), Inches(8.90), Inches(0.70))
    tf_bn = tb_bn.text_frame
    tf_bn.word_wrap = True
    tf_bn.margin_left = tf_bn.margin_top = tf_bn.margin_right = tf_bn.margin_bottom = 0
    p = tf_bn.paragraphs[0]
    p.text = "Strictly Frozen 80-Query Evaluation Benchmark (SHA-256 Verified)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_BLUE
    p_sub = tf_bn.add_paragraph()
    p_sub.text = "• Exactly 20 queries per intent class (20 MATHDS, 20 TXTSRC, 20 SEMADPT, 20 CSTRSET).\n• Immutable benchmark hashes: eval_queries.json (83ccec72...) & reference_corpus.json (0b390fdb...).\n• Objective automated evaluation: exact arithmetic truth, corpus substring presence, and strict JSON/YAML syntax."
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(8)
    p_sub.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Side-by-Side Experimental Conditions
    w_cond = Inches(4.45)
    h_cond = Inches(3.30)

    # Condition A: Baseline
    add_header_card(slide6, Inches(0.40), Inches(1.75), w_cond, h_cond, "Condition A: Direct Monolithic LLM (Baseline)", header_bg=COLOR_SLATE_DARK)
    tb_ca = slide6.shapes.add_textbox(Inches(0.55), Inches(2.20), w_cond - Inches(0.30), h_cond - Inches(0.55))
    tf_ca = tb_ca.text_frame
    tf_ca.word_wrap = True
    tf_ca.margin_left = tf_ca.margin_top = tf_ca.margin_right = tf_ca.margin_bottom = 0

    p_a = tf_ca.paragraphs[0]
    p_a.text = "Standard Production Industry Baseline:"
    p_a.font.name = FONT_FAMILY
    p_a.font.size = Pt(9)
    p_a.font.bold = True
    p_a.font.color.rgb = COLOR_NAVY_BLUE

    items_a = [
        ("Model", "Google Gemini 3.1 Flash-Lite (Cloud API)."),
        ("Execution", "All 80 queries sent directly as raw prompts."),
        ("Tool Access", "None — LLM acts as universal black-box solver."),
        ("Resource Profile", "100% cloud token generation; full network round-trip overhead on all queries."),
        ("Evaluation", "Assesses raw foundation model accuracy, latency, and token expense across all 4 domains.")
    ]
    for k, v in items_a:
        p_k = tf_ca.add_paragraph()
        p_k.space_before = Pt(3)
        p_k.text = f"• {k}: {v}"
        p_k.font.name = FONT_FAMILY
        p_k.font.size = Pt(8)
        p_k.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Condition B: Proposed System
    add_header_card(slide6, Inches(5.15), Inches(1.75), w_cond, h_cond, "Condition B: IntentFlow Routed Pipeline (Proposed)", header_bg=COLOR_NAVY_BLUE)
    tb_cb = slide6.shapes.add_textbox(Inches(5.30), Inches(2.20), w_cond - Inches(0.30), h_cond - Inches(0.55))
    tf_cb = tb_cb.text_frame
    tf_cb.word_wrap = True
    tf_cb.margin_left = tf_cb.margin_top = tf_cb.margin_right = tf_cb.margin_bottom = 0

    p_b = tf_cb.paragraphs[0]
    p_b.text = "Pre-Generation Intent-Aware Architecture:"
    p_b.font.name = FONT_FAMILY
    p_b.font.size = Pt(9)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_NAVY_BLUE

    items_b = [
        ("Classifier", "Frozen DistilBERT 66M parameter model (~41 ms)."),
        ("Dispatch", "Autonomous routing to domain-optimal engine."),
        ("Execution", "Local SymPy (math), Corpus RAG (facts), Local Schema Synthesizer, or Gemini API (creative)."),
        ("Resource Profile", "77.5% queries execute locally at $0.00 cloud cost and sub-150ms latency."),
        ("Telemetry", "Sub-millisecond latency audit and token tracking.")
    ]
    for k, v in items_b:
        p_k = tf_cb.add_paragraph()
        p_k.space_before = Pt(3)
        p_k.text = f"• {k}: {v}"
        p_k.font.name = FONT_FAMILY
        p_k.font.size = Pt(8)
        p_k.font.color.rgb = COLOR_TEXT_CHARCOAL

    set_notes(slide6, (
        "To rigorously evaluate the architecture, we conducted a matched-pair benchmark across 80 curated queries, "
        "with exactly 20 queries per intent class. "
        "In Condition A, all 80 queries were submitted directly to Google Gemini 3.1 Flash-Lite, representing standard practice. "
        "In Condition B, the exact same 80 queries were processed through our IntentFlow pipeline. "
        "We measured benchmark correctness, median and mean latency, token consumption, API calls, and total financial cost."
    ))

    # =========================================================================
    # SLIDE 7 — FINAL RESULTS (WITH EFFICIENCY COMPARISON GRAPH)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide7, "Final Results: The Cost–Accuracy–Latency Tradeoff", 7)

    # 4 Large Stat KPI Cards across top
    top_w = Inches(2.15)
    top_h = Inches(1.05)
    top_lefts = [Inches(0.40) + i * (top_w + Inches(0.20)) for i in range(4)]

    add_stat_card(slide7, top_lefts[0], Inches(0.75), top_w, top_h,
                  "BENCHMARK ACCURACY", "96.25%", "Direct: 100.0%  |  Delta: -3.75% pts (77/80)", COLOR_NAVY_BLUE)

    add_stat_card(slide7, top_lefts[1], Inches(0.75), top_w, top_h,
                  "MEDIAN LATENCY (P50)", "139 ms", "Direct: 1,922 ms  |  92.8% Faster (13.8x)", COLOR_ACCENT_BLUE)

    add_stat_card(slide7, top_lefts[2], Inches(0.75), top_w, top_h,
                  "API COST REDUCTION", "68.15%", "Direct: $0.00223  →  Routed: $0.00071", COLOR_SUCCESS_GREEN)

    add_stat_card(slide7, top_lefts[3], Inches(0.75), top_w, top_h,
                  "CLOUD API OFFLOAD", "77.50%", "80 calls reduced to 18 (62 local @ $0)", COLOR_SUCCESS_GREEN)

    # Main Visual: Efficiency Comparison Graph (Wide 3-panel figure)
    eff_img_path = PROJECT_ROOT / "artifacts" / "figures" / "fig8_efficiency_comparison.png"
    if eff_img_path.exists():
        slide7.shapes.add_picture(str(eff_img_path), Inches(0.40), Inches(1.90), width=Inches(9.20), height=Inches(2.85))

    # Bottom Core Scientific Takeaway Strip
    add_card(slide7, Inches(0.40), Inches(4.85), Inches(9.20), Inches(0.45), bg_color=COLOR_WHITE, border_color=COLOR_ACCENT_BLUE)
    tb_res_nt = slide7.shapes.add_textbox(Inches(0.55), Inches(4.88), Inches(8.90), Inches(0.35))
    tf_res_nt = tb_res_nt.text_frame
    tf_res_nt.word_wrap = True
    tf_res_nt.margin_left = tf_res_nt.margin_top = tf_res_nt.margin_right = tf_res_nt.margin_bottom = 0
    p = tf_res_nt.paragraphs[0]
    p.text = "Core Empirical Takeaway: Intent routing establishes a highly favorable Pareto frontier—slashing token usage by 67.1%, API cost by 68.2%, and median latency by 92.8%, with a modest 3.75 percentage-point delta in automated benchmark correctness."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_BLUE

    set_notes(slide7, (
        "Here are the definitive findings from our final evaluation. "
        "Direct Gemini achieved 100% accuracy, while our routed IntentFlow system achieved 96.25%—a modest 3.75 percentage-point delta. "
        "In exchange for this small variance, the system delivered transformative operational efficiencies: "
        "median latency plummeted from 1,922 milliseconds down to 139 milliseconds, a 92.8% speedup. "
        "Total token consumption dropped by 67.1%, financial cost was cut by 68.2%, and 77.5% of queries were completely offloaded from the cloud."
    ))

    # =========================================================================
    # SLIDE 8 — FAILURE ANALYSIS (WITH PER-INTENT ACCURACY GRAPH)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide8, "Per-Intent Breakdown & Failure Analysis", 8)

    # Left: Per-Intent Accuracy Graph Figure
    intent_img_path = PROJECT_ROOT / "artifacts" / "figures" / "fig9_per_intent_accuracy.png"
    if intent_img_path.exists():
        slide8.shapes.add_picture(str(intent_img_path), Inches(0.40), Inches(0.80), width=Inches(4.80), height=Inches(2.88))

    # Under-Graph Summary Card
    add_card(slide8, Inches(0.40), Inches(3.78), Inches(4.80), Inches(1.35), bg_color=COLOR_WHITE, border_color=COLOR_SUCCESS_GREEN)
    tb_ug = slide8.shapes.add_textbox(Inches(0.50), Inches(3.85), Inches(4.60), Inches(1.20))
    tf_ug = tb_ug.text_frame
    tf_ug.word_wrap = True
    tf_ug.margin_left = tf_ug.margin_top = tf_ug.margin_right = tf_ug.margin_bottom = 0
    p = tf_ug.paragraphs[0]
    p.text = "Operational Strengths by Intent Domain:"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_BLUE
    p_b1 = tf_ug.add_paragraph()
    p_b1.text = "• MATHDS: 100.0% (20/20) — Exact SymPy math; 0 arithmetic slips."
    p_b1.font.size = Pt(8)
    p_b1.font.color.rgb = COLOR_TEXT_CHARCOAL
    p_b2 = tf_ug.add_paragraph()
    p_b2.text = "• CSTRSET: 100.0% (20/20) — Perfect JSON/YAML syntax; 0 preambles."
    p_b2.font.size = Pt(8)
    p_b2.font.color.rgb = COLOR_TEXT_CHARCOAL
    p_b3 = tf_ug.add_paragraph()
    p_b3.text = "• TXTSRC: 95.0% (19/20) — Grounded passages with [REF_XX] tags."
    p_b3.font.size = Pt(8)
    p_b3.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Right: Comprehensive Audit of 3 Failures Card
    rf_x = Inches(5.40)
    rf_w = Inches(4.20)
    rf_h = Inches(4.33)
    add_header_card(slide8, rf_x, Inches(0.80), rf_w, rf_h, "Audit of 3 Routed Failures (77/80 Correct)", header_bg=COLOR_SLATE_DARK)
    tb_fail = slide8.shapes.add_textbox(rf_x + Inches(0.12), Inches(1.22), rf_w - Inches(0.24), rf_h - Inches(0.48))
    tf_fail = tb_fail.text_frame
    tf_fail.word_wrap = True
    tf_fail.margin_left = tf_fail.margin_top = tf_fail.margin_right = tf_fail.margin_bottom = 0

    fail_items = [
        ("1. Evaluator String Artifact (TXT_19)",
         "• Query: 'Who determined penicillin structure...?'\n"
         "• System Output: Retrieved reference citing 'Dorothy Crowfoot Hodgkin'.\n"
         "• Evaluation Root Cause: Regex strictly matched 'Dorothy Hodgkin'. Grounded fact was retrieved; failure is a test harness regex artifact."),

        ("2. Creative vs. Structural Boundary (SEM_08 & SEM_17)",
         "• Queries: Rhyming poem (AABB meter) and haiku (5-7-5 syllables).\n"
         "• System Output: Classifier predicted CSTRSET due to explicit constraint tokens ('AABB meter', 'haiku 5-7-5').\n"
         "• Schema handler returned schema fallback rather than creative text."),

        ("Key Scientific Insight",
         "• The primary failure mode of pre-generation routing is the SEMADPT / CSTRSET boundary dilemma.\n"
         "• Classifiers must separate creative stylistic rules from schema syntax.")
    ]

    for idx, (hd, body) in enumerate(fail_items):
        p_h = tf_fail.paragraphs[0] if idx == 0 else tf_fail.add_paragraph()
        if idx > 0: p_h.space_before = Pt(4)
        p_h.text = hd
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(8.5)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_NAVY_BLUE

        p_b = tf_fail.add_paragraph()
        p_b.text = body
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(7.5)
        p_b.font.color.rgb = COLOR_TEXT_CHARCOAL

    set_notes(slide8, (
        "Looking at the per-intent performance on the left, MATHDS and CSTRSET achieved a perfect 100% accuracy, "
        "completely eliminating arithmetic errors and broken JSON syntax. "
        "When we audited the three failures on the right, one—TXT_19—was an automated evaluator string artifact: "
        "the system retrieved 'Dorothy Crowfoot Hodgkin', but the regex searched for 'Dorothy Hodgkin'. "
        "The remaining two failures—SEM_08 and SEM_17—were creative poems with syllable and meter rules that were misrouted to CSTRSET. "
        "This identifies our key research insight: distinguishing creative constraints from structural formatting is the primary boundary dilemma."
    ))

    # =========================================================================
    # SLIDE 9 — CONCLUSION & FUTURE WORK (LOW TEXT)
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_standard_chrome(slide9, "Conclusion & Future Research Directions", 9)

    w_half = Inches(4.45)
    h_half = Inches(4.25)

    # Left: Conclusion
    add_header_card(slide9, Inches(0.40), Inches(0.85), w_half, h_half, "Conclusions & Scientific Contributions", header_bg=COLOR_NAVY_BLUE)
    tb_con = slide9.shapes.add_textbox(Inches(0.55), Inches(1.30), w_half - Inches(0.30), h_half - Inches(0.50))
    tf_con = tb_con.text_frame
    tf_con.word_wrap = True
    tf_con.margin_left = tf_con.margin_top = tf_con.margin_right = tf_con.margin_bottom = 0

    conclusions = [
        ("Intent Routing is Highly Practical",
         "A compact 66M parameter model provides accurate (<50ms) gating to specialize query processing."),

        ("Eliminates Formal Hallucinations",
         "Delegating arithmetic and schema formatting to symbolic tools guarantees 100% accuracy with zero hallucinations."),

        ("Substantial Operational Dividends",
         "Offloading 77.5% of queries slashes API expenses by 68.2%, token volume by 67.1%, and median latency by 92.8%."),

        ("Clear Pareto Frontier Characterization",
         "Validates that accepting a modest 3.75% delta in automated benchmark accuracy yields transformative cost and latency benefits.")
    ]

    for idx, (hd, body) in enumerate(conclusions):
        p_h = tf_con.paragraphs[0] if idx == 0 else tf_con.add_paragraph()
        if idx > 0: p_h.space_before = Pt(5)
        p_h.text = f"• {hd}"
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(9.5)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_NAVY_BLUE

        p_b = tf_con.add_paragraph()
        p_b.text = f"  {body}"
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(8)
        p_b.font.color.rgb = COLOR_TEXT_CHARCOAL

    # Right: Future Work
    add_header_card(slide9, Inches(5.15), Inches(0.85), w_half, h_half, "Future Research Directions", header_bg=COLOR_SLATE_DARK)
    tb_fut = slide9.shapes.add_textbox(Inches(5.30), Inches(1.30), w_half - Inches(0.30), h_half - Inches(0.50))
    tf_fut = tb_fut.text_frame
    tf_fut.word_wrap = True
    tf_fut.margin_left = tf_fut.margin_top = tf_fut.margin_right = tf_fut.margin_bottom = 0

    future_items = [
        ("Confidence-Aware Generative Fallback",
         "Automatically fall back to cloud LLMs when classifier confidence on boundary queries drops below ~85%."),

        ("Contrastive Boundary Tuning",
         "Fine-tune representations using contrastive loss to better decouple poetic meter from structured schemas."),

        ("Dynamic Runtime Schema Parsing",
         "Expand the local schema engine to dynamically parse arbitrary user schemas at runtime."),

        ("Multi-Model Routing Portfolios",
         "Extend router to dynamically select among multiple cloud LLMs (Gemini, Claude, LLaMA) based on task complexity.")
    ]

    for idx, (hd, body) in enumerate(future_items):
        p_h = tf_fut.paragraphs[0] if idx == 0 else tf_fut.add_paragraph()
        if idx > 0: p_h.space_before = Pt(5)
        p_h.text = f"{idx + 1}. {hd}"
        p_h.font.name = FONT_FAMILY
        p_h.font.size = Pt(9.5)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_NAVY_BLUE

        p_b = tf_fut.add_paragraph()
        p_b.text = f"   {body}"
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(8)
        p_b.font.color.rgb = COLOR_TEXT_CHARCOAL

    set_notes(slide9, (
        "To conclude, our research proves that pre-generation intent routing is an effective, economically viable strategy. "
        "By substituting local symbolic engines and retrieval modules for routine math, factual, and schema queries, "
        "we eliminate hallucinations in formal domains and slash operating costs by 68% and median latency by 93%. "
        "It demonstrates a practical Pareto frontier for enterprise AI deployments. "
        "For future work, we propose implementing confidence-aware fallbacks, contrastive fine-tuning for creative constraints, "
        "and multi-model routing portfolios."
    ))

    # =========================================================================
    # SLIDE 10 — THANK YOU
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    bg10_top = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(0.0), Inches(10.0), Inches(3.60))
    bg10_top.fill.solid()
    bg10_top.fill.fore_color.rgb = COLOR_NAVY_DARK
    bg10_top.line.fill.background()

    bg10_bot = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(3.60), Inches(10.0), Inches(2.025))
    bg10_bot.fill.solid()
    bg10_bot.fill.fore_color.rgb = COLOR_WHITE
    bg10_bot.line.fill.background()

    line10 = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(3.57), Inches(10.0), Inches(0.03))
    line10.fill.solid()
    line10.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    line10.line.fill.background()

    tb_ty = slide10.shapes.add_textbox(Inches(0.60), Inches(0.90), Inches(8.80), Inches(1.80))
    tf_ty = tb_ty.text_frame
    tf_ty.word_wrap = True
    tf_ty.margin_left = tf_ty.margin_top = tf_ty.margin_right = tf_ty.margin_bottom = 0

    p_ty = tf_ty.paragraphs[0]
    p_ty.text = "Thank You!"
    p_ty.font.name = FONT_FAMILY
    p_ty.font.size = Pt(38)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_WHITE

    p_qd = tf_ty.add_paragraph()
    p_qd.space_before = Pt(8)
    p_qd.text = "Questions & Discussion"
    p_qd.font.name = FONT_FAMILY
    p_qd.font.size = Pt(18)
    p_qd.font.bold = True
    p_qd.font.color.rgb = COLOR_TEXT_LIGHT

    p_proj = tf_ty.add_paragraph()
    p_proj.space_before = Pt(4)
    p_proj.text = "User Intent Classification for Hallucination Reduction in LLMs: A Cost-Accuracy Tradeoff Study"
    p_proj.font.name = FONT_FAMILY
    p_proj.font.size = Pt(11)
    p_proj.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    tb_ty_bot = slide10.shapes.add_textbox(Inches(0.60), Inches(3.85), Inches(8.80), Inches(1.20))
    tf_ty_bot = tb_ty_bot.text_frame
    tf_ty_bot.word_wrap = True
    tf_ty_bot.margin_left = tf_ty_bot.margin_top = tf_ty_bot.margin_right = tf_ty_bot.margin_bottom = 0

    p_pres = tf_ty_bot.paragraphs[0]
    p_pres.text = "P.U. Graandhikaa Sri (CH.SC.U4CSE23136)     T. L. Bharath (CH.SC.U4CSE23149)"
    p_pres.font.name = FONT_FAMILY
    p_pres.font.size = Pt(13)
    p_pres.font.bold = True
    p_pres.font.color.rgb = COLOR_TEXT_DARK

    p_res = tf_ty_bot.add_paragraph()
    p_res.space_before = Pt(4)
    p_res.text = "Research Artifacts & Project Resources:"
    p_res.font.name = FONT_FAMILY
    p_res.font.size = Pt(9.5)
    p_res.font.bold = True
    p_res.font.color.rgb = COLOR_NAVY_BLUE

    items_res = [
        "• Master Technical Report: artifacts/reports/MASTER_TECHNICAL_RESEARCH_REPORT.docx (.pdf, 34 pages)",
        "• Research Record: docs/research_record/PROJECT_RESEARCH_RECORD.md (Phases 1 through 6.8 & Literature Review)",
        "• Interactive Demonstration: IntentFlow Web Application (Flask, Local CPU Inference at http://127.0.0.1:5000)"
    ]
    for it in items_res:
        p_it = tf_ty_bot.add_paragraph()
        p_it.text = it
        p_it.font.name = FONT_FAMILY
        p_it.font.size = Pt(8)
        p_it.font.color.rgb = COLOR_TEXT_MUTED

    set_notes(slide10, (
        "Thank you for your time and attention. "
        "All model checkpoints, datasets, cryptographic benchmarks, and demonstration code are fully documented. "
        "We are now glad to take any questions or feedback from the committee."
    ))

    # Save presentation
    output_path = PROJECT_ROOT / "artifacts" / "reports" / "FINAL_PROJECT_PRESENTATION.pptx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"[SUCCESS] Saved 10-slide presentation to: {output_path}")

if __name__ == "__main__":
    create_deck()
