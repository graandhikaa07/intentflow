"""
Master Technical Research Report Generator
Generates:
1. artifacts/reports/MASTER_TECHNICAL_RESEARCH_REPORT.docx
2. artifacts/reports/MASTER_TECHNICAL_RESEARCH_REPORT.pdf (via MS Word COM)
"""

import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT_ROOT = Path(__file__).resolve().parent.parent

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex.replace('#', ''))
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell margins in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    Sets cell borders.
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='CBD5E1')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_callout(doc, text, title=None, border_color="0284C7", bg_color="F8FAFC"):
    """Adds a stylish callout block with a thick left border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=140)
    set_cell_border(cell, 
                    left=dict(val='single', sz=24, color=border_color),
                    top=dict(val='single', sz=4, color='E2E8F0'),
                    bottom=dict(val='single', sz=4, color='E2E8F0'),
                    right=dict(val='single', sz=4, color='E2E8F0'))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_t = p.add_run(f"{title}\n")
        run_t.bold = True
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(10.5)
        run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    run_txt = p.add_run(text)
    run_txt.font.name = 'Calibri'
    run_txt.font.size = Pt(9.5)
    run_txt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Empty space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

print("Helper functions defined successfully.")
