import pptx
from pptx import Presentation
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation('docs/Review2_UserIntent_LLM.pptx')

def inspect_slide(slide_idx):
    slide = prs.slides[slide_idx]
    print("=" * 80)
    print(f"SLIDE {slide_idx + 1}")
    # Slide background
    print(f"Background: {slide.background}")
    for s_idx, shape in enumerate(slide.shapes):
        left, top, w, h = shape.left/914400, shape.top/914400, shape.width/914400, shape.height/914400
        fill_str = "None"
        try:
            if shape.fill.type:
                color = shape.fill.fore_color.rgb if hasattr(shape.fill.fore_color, 'rgb') else 'theme'
                fill_str = f"{shape.fill.type} ({color})"
        except:
            pass
        line_str = "None"
        try:
            if shape.line.fill.type:
                color = shape.line.color.rgb if hasattr(shape.line.color, 'rgb') else 'theme'
                line_str = f"{shape.line.fill.type} ({color})"
        except:
            pass
        print(f"  Shape {s_idx} [{shape.name}] {shape.shape_type} pos=({left:.2f}, {top:.2f}) size=({w:.2f} x {h:.2f}) fill={fill_str} line={line_str}")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    font_name = p.font.name if p.font else None
                    font_size = p.font.size.pt if (p.font and p.font.size) else None
                    bold = p.font.bold if p.font else None
                    color = None
                    try:
                        color = p.font.color.rgb if (p.font and p.font.color and hasattr(p.font.color, 'rgb')) else None
                    except:
                        pass
                    print(f"    Text: [font={font_name}, size={font_size}, bold={bold}, color={color}] \"{txt}\"")

inspect_slide(0)  # Slide 1
inspect_slide(1)  # Slide 2
inspect_slide(7)  # Slide 8
inspect_slide(13) # Slide 14
