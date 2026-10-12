import pptx
from pptx import Presentation
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation('docs/Review2_UserIntent_LLM.pptx')

for s_idx in [0, 1, 2, 7]:
    slide = prs.slides[s_idx]
    print(f"=== Slide {s_idx+1} ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    txt = r.text.strip()
                    if txt:
                        fn = r.font.name
                        fs = r.font.size.pt if r.font.size else None
                        b = r.font.bold
                        c = None
                        try:
                            c = r.font.color.rgb if (r.font.color and hasattr(r.font.color, 'rgb')) else None
                        except:
                            pass
                        print(f"  Run: {txt[:40]!r} | font={fn} | size={fs} | bold={b} | color={c}")
